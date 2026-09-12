import io
import os
import re
import secrets
import time
import asyncio
from typing import Any, Literal
from zipfile import BadZipFile

from docx import Document
from fastapi import FastAPI, File, Form, Header, HTTPException, UploadFile
from pydantic import BaseModel, Field
from pypdf import PdfReader

TOKEN = os.getenv("DOCLING_INTERNAL_TOKEN", "")
MAX_BYTES = int(os.getenv("DOCLING_MAX_INPUT_BYTES", str(10 * 1024 * 1024)))
MAX_BLOCKS = int(os.getenv("DOCLING_MAX_BLOCKS", "1500"))
MAX_CONCURRENCY = max(1, int(os.getenv("DOCLING_MAX_CONCURRENCY", "2")))
_slots = asyncio.Semaphore(MAX_CONCURRENCY)

app = FastAPI(title="RoleBrief private Docling service", docs_url=None, redoc_url=None, openapi_url=None)


class BoundingBox(BaseModel):
    x: float
    y: float
    width: float
    height: float


class Page(BaseModel):
    pageNumber: int
    width: float | None = None
    height: float | None = None


class Block(BaseModel):
    id: str
    kind: str
    text: str = Field(max_length=8000)
    pageNumber: int | None
    boundingBox: BoundingBox | None = None
    readingOrder: int


class WarningItem(BaseModel):
    code: str
    message: str | None = None


class Metrics(BaseModel):
    pageCount: int
    blockCount: int
    ocrUsed: bool
    processingMs: int


class ExtractionResponse(BaseModel):
    contractVersion: Literal[1]
    parserName: Literal["docling"]
    parserVersion: str
    document: dict[str, Any]
    pages: list[Page]
    blocks: list[Block]
    warnings: list[WarningItem]
    metrics: Metrics


@app.get("/internal/v1/health/live")
def live():
    return {"status": "ok"}


@app.get("/internal/v1/health/ready")
def ready():
    return {"status": "ok"}


@app.post("/internal/v1/documents/extract", response_model=ExtractionResponse)
async def extract(
    contractVersion: int = Form(...),
    documentRef: str = Form(...),
    mediaType: str = Form(...),
    ocrPolicy: str = Form(...),
    deadlineMs: int = Form(...),
    file: UploadFile = File(...),
    authorization: str | None = Header(default=None),
    x_rolebrief_contract: str | None = Header(default=None),
):
    if not TOKEN or not authorization or not secrets.compare_digest(authorization, f"Bearer {TOKEN}"):
        raise HTTPException(status_code=401, detail="unauthorized")
    if contractVersion != 1 or x_rolebrief_contract != "docling-v1":
        raise HTTPException(status_code=400, detail="invalid_contract")
    if not re.fullmatch(r"resume:[A-Za-z0-9_-]+", documentRef):
        raise HTTPException(status_code=400, detail="invalid_document_ref")
    if _slots.locked() and getattr(_slots, "_value", 0) <= 0:
        raise HTTPException(status_code=503, detail="docling_overloaded", headers={"Retry-After": "30"})
    async with _slots:
        return await extract_bounded(contractVersion, documentRef, mediaType, ocrPolicy, deadlineMs, file)


async def extract_bounded(
    contractVersion: int,
    documentRef: str,
    mediaType: str,
    ocrPolicy: str,
    deadlineMs: int,
    file: UploadFile,
):
    started = time.perf_counter()
    data = await file.read(MAX_BYTES + 1)
    if len(data) == 0:
        raise HTTPException(status_code=422, detail="empty_document")
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="document_too_large")
    if mediaType == "application/pdf":
        pages, blocks, warnings = extract_pdf(data)
    elif mediaType == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        pages, blocks, warnings = extract_docx(data)
    else:
        raise HTTPException(status_code=415, detail="unsupported_media_type")
    if len(blocks) > MAX_BLOCKS:
        raise HTTPException(status_code=413, detail="too_many_blocks")
    usable = sum(len(block.text.strip()) for block in blocks)
    if usable == 0:
        code = "DOCX_IMAGE_ONLY_UNSUPPORTED" if mediaType.endswith("document") else "NO_USABLE_CONTENT"
        raise HTTPException(status_code=422, detail=code)
    elapsed = int((time.perf_counter() - started) * 1000)
    return ExtractionResponse(
        contractVersion=1,
        parserName="docling",
        parserVersion="rolebrief-docling-service-v1",
        document={"documentRef": documentRef, "mediaType": mediaType},
        pages=pages,
        blocks=blocks,
        warnings=warnings,
        metrics=Metrics(pageCount=len(pages), blockCount=len(blocks), ocrUsed=False, processingMs=elapsed),
    )


def extract_pdf(data: bytes):
    warnings: list[WarningItem] = []
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise HTTPException(status_code=422, detail="PDF_ENCRYPTED")
        pages = [
            Page(pageNumber=i + 1, width=float(page.mediabox.width), height=float(page.mediabox.height))
            for i, page in enumerate(reader.pages)
        ]
        blocks: list[Block] = []
        order = 0
        for page_index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            for paragraph in split_blocks(text):
                blocks.append(Block(id=f"p{page_index + 1}_b{order}", kind=kind_for(paragraph), text=paragraph, pageNumber=page_index + 1, readingOrder=order))
                order += 1
        if not blocks:
            warnings.append(WarningItem(code="OCR_NOT_AVAILABLE_IN_LOCAL_SERVICE", message="No text extracted from PDF."))
        return pages, blocks, warnings
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="PDF_INVALID")


def extract_docx(data: bytes):
    try:
        doc = Document(io.BytesIO(data))
    except BadZipFile:
        raise HTTPException(status_code=422, detail="DOCX_INVALID")
    except Exception:
        raise HTTPException(status_code=422, detail="DOCX_INVALID")
    blocks: list[Block] = []
    order = 0
    for para in doc.paragraphs:
        text = clean(para.text)
        if not text:
            continue
        style = (para.style.name if para.style else "").lower()
        kind = "heading" if "heading" in style or len(text) <= 40 and text.upper() == text else kind_for(text)
        blocks.append(Block(id=f"d_b{order}", kind=kind, text=text, pageNumber=None, readingOrder=order))
        order += 1
    for table in doc.tables:
        rows = []
        for row in table.rows[:50]:
            rows.append(" | ".join(clean(cell.text) for cell in row.cells[:10]))
        text = "\n".join(row for row in rows if row.strip())
        if text:
            blocks.append(Block(id=f"d_b{order}", kind="table", text=text[:8000], pageNumber=None, readingOrder=order))
            order += 1
    return [Page(pageNumber=1)], blocks, []


def split_blocks(text: str):
    return [clean(part) for part in re.split(r"\n{2,}|\r\n{2,}", text) if clean(part)]


def clean(text: str):
    return re.sub(r"\s+", " ", text).strip()


def kind_for(text: str):
    normalized = clean(text)
    if len(normalized) <= 40 and not normalized.endswith("."):
        return "heading"
    if " | " in normalized:
        return "table"
    return "paragraph"
