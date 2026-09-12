import asyncio
import io
import os
import re
import secrets
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile

import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from fastapi import FastAPI, File, Form, Header, HTTPException, UploadFile
from pydantic import BaseModel, Field

SERVICE_VERSION = "rolebrief-document-parser-service-v1"
TOKEN = os.getenv("DOCUMENT_PARSER_INTERNAL_TOKEN", "")
MAX_BYTES = int(os.getenv("DOCUMENT_PARSER_MAX_INPUT_BYTES", str(10 * 1024 * 1024)))
MAX_BLOCKS = int(os.getenv("DOCUMENT_PARSER_MAX_BLOCKS", "1500"))
MAX_CHARS = int(os.getenv("DOCUMENT_PARSER_MAX_CHARS", str(2 * 1024 * 1024)))
MAX_BLOCK_CHARS = int(os.getenv("DOCUMENT_PARSER_MAX_BLOCK_CHARS", "8000"))
OCR_MIN_CHARS_PER_PAGE = int(os.getenv("DOCUMENT_PARSER_OCR_MIN_CHARS_PER_PAGE", "80"))
MAX_CONCURRENCY = max(1, int(os.getenv("DOCUMENT_PARSER_MAX_CONCURRENCY", "2")))
TESSERACT_TIMEOUT_SECONDS = max(1, int(os.getenv("DOCUMENT_PARSER_TESSERACT_TIMEOUT_SECONDS", "15")))
TESSERACT_LANG = os.getenv("DOCUMENT_PARSER_TESSERACT_LANG", "eng")
_slots = asyncio.Semaphore(MAX_CONCURRENCY)

app = FastAPI(title="RoleBrief private document parser service", docs_url=None, redoc_url=None, openapi_url=None)


class BoundingBox(BaseModel):
    x0: float
    y0: float
    x1: float
    y1: float


class Page(BaseModel):
    pageNumber: int
    width: float | None = None
    height: float | None = None


class Block(BaseModel):
    id: str
    kind: Literal["HEADING", "PARAGRAPH", "LIST_ITEM", "TABLE", "LINK", "UNKNOWN"]
    text: str = Field(max_length=8000)
    pageNumber: int | None
    boundingBox: BoundingBox | None = None
    readingOrder: int
    extractionMethod: Literal["NATIVE_PDF", "DOCX_XML", "OCR"]


class WarningItem(BaseModel):
    code: str
    pageNumber: int | None = None


class ParserMetadata(BaseModel):
    serviceVersion: str
    nativePdfEngine: str | None
    docxEngine: str | None
    ocrEngine: str | None
    ocrEngineVersion: str | None


class Metrics(BaseModel):
    pageCount: int
    blockCount: int
    characterCount: int
    ocrPages: int
    processingMs: int


class ExtractionResponse(BaseModel):
    contractVersion: Literal[1]
    parser: ParserMetadata
    pages: list[Page]
    blocks: list[Block]
    warnings: list[WarningItem]
    metrics: Metrics


@app.get("/internal/v1/health/live")
def live():
    return {"status": "ok"}


@app.get("/internal/v1/health/ready")
def ready():
    return {
        "status": "ok",
        "serviceVersion": SERVICE_VERSION,
        "nativePdfEngine": "pdfplumber",
        "docxEngine": "python-docx",
        "ocrEngine": "tesseract",
        "ocrEngineVersion": tesseract_version(),
        "maxConcurrency": MAX_CONCURRENCY
    }


@app.post("/internal/v1/documents/extract", response_model=ExtractionResponse)
async def extract(
    contractVersion: int = Form(...),
    documentRef: str = Form(...),
    mediaType: str = Form(...),
    ocrPolicy: Literal["auto", "disabled"] = Form("auto"),
    deadlineMs: int = Form(30000),
    file: UploadFile = File(...),
    authorization: str | None = Header(default=None),
    x_rolebrief_contract: str | None = Header(default=None),
):
    if TOKEN and not secrets.compare_digest(authorization or "", f"Bearer {TOKEN}"):
        raise HTTPException(status_code=401, detail="unauthorized")
    if contractVersion != 1 or x_rolebrief_contract != "document-parser-v1":
        raise HTTPException(status_code=400, detail="invalid_contract")
    if not re.fullmatch(r"resume:[A-Za-z0-9_-]+", documentRef):
        raise HTTPException(status_code=400, detail="invalid_document_ref")
    if mediaType not in {"application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}:
        raise HTTPException(status_code=415, detail="unsupported_media_type")
    if _slots.locked() and getattr(_slots, "_value", 0) <= 0:
        raise HTTPException(status_code=503, detail="parser_overloaded", headers={"Retry-After": "30"})
    async with _slots:
        try:
            return await asyncio.wait_for(extract_bounded(mediaType, ocrPolicy, file), timeout=max(1, deadlineMs / 1000))
        except asyncio.TimeoutError as exc:
            raise HTTPException(status_code=504, detail="PARSER_TIMEOUT") from exc


async def extract_bounded(media_type: str, ocr_policy: str, file: UploadFile):
    started = time.perf_counter()
    data = await file.read(MAX_BYTES + 1)
    if not data:
        raise HTTPException(status_code=422, detail="NO_USABLE_CONTENT")
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="PARSER_OUTPUT_TOO_LARGE")

    if media_type == "application/pdf":
        pages, blocks, warnings, ocr_pages = await asyncio.to_thread(extract_pdf, data, ocr_policy)
        parser = ParserMetadata(
            serviceVersion=SERVICE_VERSION,
            nativePdfEngine="pdfplumber/pdfminer.six",
            docxEngine=None,
            ocrEngine="tesseract" if ocr_pages else None,
            ocrEngineVersion=tesseract_version() if ocr_pages else None
        )
    else:
        pages, blocks, warnings = await asyncio.to_thread(extract_docx, data)
        ocr_pages = 0
        parser = ParserMetadata(
            serviceVersion=SERVICE_VERSION,
            nativePdfEngine=None,
            docxEngine="python-docx",
            ocrEngine=None,
            ocrEngineVersion=None
        )

    character_count = sum(len(block.text) for block in blocks)
    if len(blocks) > MAX_BLOCKS or character_count > MAX_CHARS:
        raise HTTPException(status_code=413, detail="PARSER_OUTPUT_TOO_LARGE")
    if character_count == 0:
        raise HTTPException(status_code=422, detail="NO_USABLE_CONTENT")

    return ExtractionResponse(
        contractVersion=1,
        parser=parser,
        pages=pages,
        blocks=blocks,
        warnings=warnings,
        metrics=Metrics(
            pageCount=len(pages),
            blockCount=len(blocks),
            characterCount=character_count,
            ocrPages=ocr_pages,
            processingMs=int((time.perf_counter() - started) * 1000)
        )
    )


def extract_pdf(data: bytes, ocr_policy: str):
    pages: list[Page] = []
    blocks: list[Block] = []
    warnings: list[WarningItem] = []
    ocr_pages = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                pages.append(Page(pageNumber=page_index, width=float(page.width), height=float(page.height)))
                page_blocks = blocks_from_pdf_page(page, page_index, len(blocks))
                if should_ocr(page_blocks) and ocr_policy == "auto":
                    try:
                        ocr_blocks = ocr_pdf_page(data, page_index, len(blocks))
                        if ocr_blocks:
                            blocks.extend(ocr_blocks)
                            ocr_pages += 1
                            continue
                    except Exception:
                        warnings.append(WarningItem(code="OCR_FAILED", pageNumber=page_index))
                if uncertain_columns(page):
                    warnings.append(WarningItem(code="PDF_READING_ORDER_UNCERTAIN", pageNumber=page_index))
                blocks.extend(page_blocks)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="PDF_EXTRACTION_FAILED") from exc
    return pages, blocks, warnings, ocr_pages


def blocks_from_pdf_page(page, page_number: int, start_order: int):
    words = page.extract_words(use_text_flow=False, keep_blank_chars=False) or []
    ordered_words = sorted(words, key=lambda word: float(word["x0"]))
    gaps = [
        (float(ordered_words[index + 1]["x0"]) - float(ordered_words[index]["x1"]), index)
        for index in range(len(ordered_words) - 1)
    ]
    column_groups: list[list[dict]]
    if gaps:
        gap, index = max(gaps)
        if gap >= 80:
            column_groups = [ordered_words[: index + 1], ordered_words[index + 1:]]
        else:
            column_groups = [words]
    else:
        column_groups = [words]

    ordered_lines: list[list[dict]] = []
    for column in column_groups:
        lines: dict[int, list[dict]] = {}
        for word in column:
            bucket = round(float(word["top"]) / 4)
            lines.setdefault(bucket, []).append(word)
        ordered_lines.extend(
            sorted(lines.values(), key=lambda row: (min(float(w["top"]) for w in row), min(float(w["x0"]) for w in row)))
        )
    blocks: list[Block] = []
    for offset, row in enumerate(ordered_lines):
        row = sorted(row, key=lambda word: float(word["x0"]))
        text = clean_text(" ".join(word["text"] for word in row))
        if not text:
            continue
        x0 = min(float(word["x0"]) for word in row)
        y0 = min(float(word["top"]) for word in row)
        x1 = max(float(word["x1"]) for word in row)
        y1 = max(float(word["bottom"]) for word in row)
        blocks.append(Block(
            id=f"p{page_number}_b{start_order + offset}",
            kind=kind_for(text),
            text=text[:MAX_BLOCK_CHARS],
            pageNumber=page_number,
            boundingBox=BoundingBox(x0=x0, y0=y0, x1=x1, y1=y1),
            readingOrder=start_order + offset,
            extractionMethod="NATIVE_PDF"
        ))
    return order_columns(blocks)


def order_columns(blocks: list[Block]):
    if len(blocks) < 4:
        return blocks
    centers = sorted(
        (((block.boundingBox.x0 + block.boundingBox.x1) / 2, block) for block in blocks if block.boundingBox),
        key=lambda item: item[0],
    )
    gaps = [(centers[i + 1][0] - centers[i][0], i) for i in range(len(centers) - 1)]
    if not gaps:
        return blocks
    gap, index = max(gaps)
    if gap < 80:
        return blocks
    left = [item[1] for item in centers[: index + 1]]
    right = [item[1] for item in centers[index + 1:]]
    ordered = sorted(left, key=lambda block: (block.boundingBox.y0, block.boundingBox.x0)) + sorted(right, key=lambda block: (block.boundingBox.y0, block.boundingBox.x0))
    for order, block in enumerate(ordered):
        block.readingOrder = order
    return ordered


def uncertain_columns(page):
    words = page.extract_words(use_text_flow=False, keep_blank_chars=False) or []
    if len(words) < 30:
        return False
    xs = sorted(float(word["x0"]) for word in words)
    gaps = [xs[i + 1] - xs[i] for i in range(len(xs) - 1)]
    return bool(gaps and max(gaps) > 100)


def should_ocr(blocks: list[Block]):
    return sum(len(block.text) for block in blocks) < OCR_MIN_CHARS_PER_PAGE


def ocr_pdf_page(data: bytes, page_number: int, start_order: int):
    tmp_pdf = None
    tmp_png = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as pdf_file:
            pdf_file.write(data)
            tmp_pdf = Path(pdf_file.name)
        doc = pdfium.PdfDocument(str(tmp_pdf))
        page = doc[page_number - 1]
        bitmap = page.render(scale=2).to_pil()
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as png_file:
            tmp_png = Path(png_file.name)
            bitmap.save(png_file, format="PNG")
        proc = subprocess.run(
            ["tesseract", str(tmp_png), "stdout", "-l", TESSERACT_LANG],
            capture_output=True,
            text=True,
            timeout=TESSERACT_TIMEOUT_SECONDS,
            check=False
        )
        if proc.returncode != 0:
            raise RuntimeError("tesseract failed")
        blocks = []
        for index, text in enumerate(split_text_blocks(proc.stdout)):
            blocks.append(Block(
                id=f"p{page_number}_ocr{start_order + index}",
                kind=kind_for(text),
                text=text[:MAX_BLOCK_CHARS],
                pageNumber=page_number,
                boundingBox=None,
                readingOrder=start_order + index,
                extractionMethod="OCR"
            ))
        return blocks
    finally:
        for path in (tmp_pdf, tmp_png):
            if path:
                path.unlink(missing_ok=True)


def extract_docx(data: bytes):
    try:
        doc = Document(io.BytesIO(data))
    except BadZipFile as exc:
        raise HTTPException(status_code=422, detail="DOCX_EXTRACTION_FAILED") from exc
    except Exception as exc:
        raise HTTPException(status_code=422, detail="DOCX_EXTRACTION_FAILED") from exc

    blocks: list[Block] = []
    warnings: list[WarningItem] = []
    order = 0
    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "").lower()
        blocks.append(Block(
            id=f"d_b{order}",
            kind="HEADING" if "heading" in style else kind_for(text),
            text=text[:MAX_BLOCK_CHARS],
            pageNumber=None,
            boundingBox=None,
            readingOrder=order,
            extractionMethod="DOCX_XML"
        ))
        order += 1
    for table in doc.tables:
        rows = []
        for row in table.rows[:50]:
            rows.append(" | ".join(clean_text(cell.text) for cell in row.cells[:10]))
        text = "\n".join(row for row in rows if row.strip())
        if text:
            blocks.append(Block(id=f"d_b{order}", kind="TABLE", text=text[:MAX_BLOCK_CHARS], pageNumber=None, boundingBox=None, readingOrder=order, extractionMethod="DOCX_XML"))
            order += 1
    if not blocks and doc.inline_shapes:
        raise HTTPException(status_code=422, detail="DOCX_IMAGE_ONLY_UNSUPPORTED")
    if not blocks:
        warnings.append(WarningItem(code="NO_USABLE_CONTENT", pageNumber=None))
    return [Page(pageNumber=1, width=None, height=None)], blocks, warnings


def split_text_blocks(text: str):
    return [clean_text(part) for part in re.split(r"\n{2,}|\r\n{2,}", text) if clean_text(part)]


def clean_text(text: str):
    return re.sub(r"\s+", " ", text).strip()


def kind_for(text: str):
    normalized = clean_text(text)
    if re.match(r"^[-*•]\s+", normalized):
        return "LIST_ITEM"
    if "|" in normalized:
        return "TABLE"
    if len(normalized) <= 80 and not normalized.endswith("."):
        return "HEADING"
    return "PARAGRAPH"


def tesseract_version():
    try:
        proc = subprocess.run(["tesseract", "--version"], capture_output=True, text=True, timeout=2, check=False)
        return (proc.stdout.splitlines()[0] if proc.stdout else "tesseract").strip()[:80]
    except Exception:
        return None
